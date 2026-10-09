"""Independent, configuration-driven postcondition checks for XLSX reports.

This module reads workbooks only. It never evaluates formulas, modifies the
artifact, imports application code, or interprets input as Python expressions.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class ContractError(ValueError):
    """The contract is invalid or cannot be applied."""


def _digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def _number(value: Any) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise ValueError("not a numeric cell")
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError("not a numeric cell") from exc
    if not result.is_finite():
        raise ValueError("non-finite numeric cell")
    return result


def _validate_contract(contract: Any) -> None:
    if not isinstance(contract, dict) or contract.get("schema_version") != 1:
        raise ContractError("schema_version must be integer 1")
    sheets = contract.get("required_sheets")
    if not isinstance(sheets, dict) or not sheets:
        raise ContractError("required_sheets must be a non-empty object")
    for name, headers in sheets.items():
        if not isinstance(name, str) or not name or not isinstance(headers, list) or not all(
            isinstance(h, str) and h for h in headers
        ) or len(headers) != len(set(headers)):
            raise ContractError("each sheet needs a list of distinct header strings")
    summary = contract.get("summary")
    if not isinstance(summary, dict) or not all(
        isinstance(summary.get(key), str) and summary[key] for key in ("sheet", "key_column", "value_column")
    ):
        raise ContractError("summary requires sheet, key_column and value_column")
    if summary["sheet"] not in sheets:
        raise ContractError("summary sheet must be in required_sheets")
    if not isinstance(contract.get("forbid_formulas", False), bool):
        raise ContractError("forbid_formulas must be boolean")
    for group, fields in (
        ("row_count_checks", ("sheet", "summary_key")),
        ("sum_checks", ("sheet", "column", "summary_key")),
    ):
        checks = contract.get(group, [])
        if not isinstance(checks, list):
            raise ContractError(f"{group} must be a list")
        for check in checks:
            if not isinstance(check, dict) or not all(isinstance(check.get(k), str) and check[k] for k in fields):
                raise ContractError(f"{group} entries require: {', '.join(fields)}")
            if check["sheet"] not in sheets:
                raise ContractError(f"{group}: sheet not in required_sheets")
            if group == "sum_checks":
                try:
                    tolerance = _number(check.get("tolerance", "0"))
                except ValueError as exc:
                    raise ContractError("sum tolerance must be finite and numeric") from exc
                if tolerance < 0:
                    raise ContractError("sum tolerance cannot be negative")


def audit(workbook_path: Path, contract: dict[str, Any]) -> dict[str, Any]:
    """Return JSON-serializable findings. PASS does not imply business approval."""
    _validate_contract(contract)
    path = Path(workbook_path)
    if not path.is_file():
        raise ContractError("workbook does not exist")
    artifact_hash = _digest(path)
    findings: list[dict[str, str]] = []

    def fail(code: str, location: str, detail: str) -> None:
        findings.append({"code": code, "location": location, "detail": detail})

    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise ContractError(f"cannot open XLSX: {type(exc).__name__}") from exc
    try:
        rows_by_sheet: dict[str, list[dict[str, Any]]] = {}
        for name, required_headers in contract["required_sheets"].items():
            if name not in workbook.sheetnames:
                fail("MISSING_SHEET", name, "required worksheet not found")
                continue
            sheet = workbook[name]
            all_rows = sheet.iter_rows(values_only=True)
            headers = list(next(all_rows, ()))
            if len(headers) != len(set(headers)):
                fail("DUPLICATE_HEADER", name, "duplicate header in first row")
            missing = [header for header in required_headers if header not in headers]
            for header in missing:
                fail("MISSING_COLUMN", name, header)
            if missing or len(headers) != len(set(headers)):
                continue
            rows_by_sheet[name] = [
                dict(zip(headers, values)) for values in all_rows
                if any(value is not None for value in values)
            ]

        if contract.get("forbid_formulas", False):
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.data_type == "f":
                            fail("FORMULA_CELL", f"{sheet.title}!{cell.coordinate}", "formula found")

        summary_conf = contract["summary"]
        summary_name = summary_conf["sheet"]
        key_col = summary_conf["key_column"]
        val_col = summary_conf["value_column"]
        metrics: dict[str, Any] = {}
        if summary_name in rows_by_sheet:
            for row in rows_by_sheet[summary_name]:
                key = row.get(key_col)
                if key is None:
                    continue
                if key in metrics:
                    fail("DUPLICATE_METRIC", summary_name, "duplicate summary key")
                metrics[str(key)] = row.get(val_col)

        for rule in contract.get("row_count_checks", []):
            name, key = rule["sheet"], rule["summary_key"]
            if name not in rows_by_sheet:
                continue
            if key not in metrics:
                fail("MISSING_METRIC", summary_name, key)
                continue
            try:
                expected = _number(metrics[key])
            except ValueError:
                fail("INVALID_METRIC", summary_name, key)
                continue
            actual = len(rows_by_sheet[name])
            if expected != actual:
                fail("COUNT_MISMATCH", name, f"summary key {key!r}: expected {expected}, observed {actual}")

        for rule in contract.get("sum_checks", []):
            name, key, col = rule["sheet"], rule["summary_key"], rule["column"]
            if name not in rows_by_sheet:
                continue
            if key not in metrics:
                fail("MISSING_METRIC", summary_name, key)
                continue
            try:
                expected = _number(metrics[key])
                tolerance = _number(rule.get("tolerance", "0"))
            except ValueError:
                fail("INVALID_METRIC", summary_name, key)
                continue
            total = Decimal("0")
            invalid = False
            for index, row in enumerate(rows_by_sheet[name], start=2):
                try:
                    total += _number(row.get(col))
                except ValueError:
                    fail("INVALID_SUM_CELL", f"{name}!{col}:{index}", "non-numeric value")
                    invalid = True
            if not invalid and abs(total - expected) > tolerance:
                fail("SUM_MISMATCH", name, f"summary key {key!r}: expected {expected}, observed {total}")
    finally:
        workbook.close()

    return {
        "schema_version": 1,
        "status": "FAIL" if findings else "PASS",
        "artifact_sha256": artifact_hash,
        "findings": findings,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Read-only contract audit of a saved XLSX report")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--contract", required=True, type=Path)
    parser.add_argument("--output", type=Path, help="Optional JSON result path")
    args = parser.parse_args(argv)
    try:
        with args.contract.open("r", encoding="utf-8") as source:
            contract = json.load(source)
        report = audit(args.workbook, contract)
        exit_code = 0 if report["status"] == "PASS" else 1
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        report = {"schema_version": 1, "status": "BLOCKED", "findings": [
            {"code": "INPUT_ERROR", "location": "contract or workbook", "detail": str(exc)}
        ]}
        exit_code = 2
    result = json.dumps(report, indent=2, sort_keys=True, ensure_ascii=False) + "\n"
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(result, encoding="utf-8")
    print(result, end="")
    return exit_code


if __name__ == "__main__":
    sys.exit(main())
