from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from generate_expense_report import (  # noqa: E402
    ORIGINAL_COLUMNS,
    REQUIRED_COLUMNS,
    parse_expense_filename,
    run_report,
    validate_and_normalize_row,
)

HEADER = ",".join(REQUIRED_COLUMNS)
EXPECTED_SHEETS = [
    "Summary",
    "By_Office",
    "By_Department",
    "By_ExpenseType",
    "Raw_Data",
    "Error_Report",
    "Run_Log",
]


def row(**overrides: str) -> dict[str, str]:
    base = {
        "Date": "2026-01-15",
        "Office": "Singapore",
        "Department": "IT",
        "Employee": "Alex Tan",
        "ExpenseType": "Meal",
        "Amount": "100.00",
        "Currency": "USD",
        "Description": "Team lunch",
    }
    return {**base, **overrides}


def write_csv(path: Path, rows: list[dict[str, str]], header: str = HEADER) -> Path:
    lines = [header]
    for r in rows:
        lines.append(",".join(f'"{r[c]}"' for c in REQUIRED_COLUMNS))
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8-sig")
    return path


def sheet_dicts(workbook, name: str) -> list[dict[str, object]]:
    rows = list(workbook[name].iter_rows(values_only=True))
    headers = rows[0]
    return [dict(zip(headers, values)) for values in rows[1:]]


def metrics(workbook) -> dict[str, object]:
    return {r["Metric"]: r["Value"] for r in sheet_dicts(workbook, "Summary")}


def log_text(output_dir: Path) -> str:
    return (output_dir / "process_log.txt").read_text(encoding="utf-8")


# --- report month validation -------------------------------------------------

def test_date_outside_report_month_is_flagged() -> None:
    result = validate_and_normalize_row(row(Date="2026-02-01"), "Singapore", "2026-01")

    assert result["IsValid"] is False
    assert result["ErrorType"] == "Mismatched Report Month"


def test_date_inside_report_month_is_valid_including_month_boundaries() -> None:
    for day in ("2026-01-01", "2026-01-31"):
        assert validate_and_normalize_row(row(Date=day), "Singapore", "2026-01")["IsValid"] is True


def test_month_check_is_skipped_when_expected_month_not_given() -> None:
    # Backward compatibility for direct callers of the validator.
    assert validate_and_normalize_row(row(Date="2025-12-31"), "Singapore")["IsValid"] is True


def test_invalid_date_does_not_also_report_month_mismatch() -> None:
    result = validate_and_normalize_row(row(Date="2026-13-45"), "Singapore", "2026-01")

    assert result["ErrorType"] == "Invalid Date"


def test_mismatched_month_row_excluded_from_totals_end_to_end(tmp_path: Path) -> None:
    write_csv(
        tmp_path / "in" / "expense_singapore_2026-01.csv",
        [row(Amount="100.00"), row(Amount="50.00", Date="2026-02-03"), row(Amount="25.50", Date="2025-12-31")],
    )
    out = tmp_path / "out"

    assert run_report(tmp_path / "in", out) == 0

    wb = load_workbook(out / "Expense_Report_2026-01.xlsx")
    m = metrics(wb)
    assert (m["Total Records"], m["Valid Records"], m["Error Records"]) == (3, 1, 2)
    assert m["Total Amount"] == 100.0
    errors = sheet_dicts(wb, "Error_Report")
    assert [e["ErrorType"] for e in errors] == ["Mismatched Report Month"] * 2
    assert [e["Original_Date"] for e in errors] == ["2026-02-03", "2025-12-31"]


# --- original value preservation --------------------------------------------

def test_original_values_preserved_separately_from_normalized() -> None:
    result = validate_and_normalize_row(
        row(Office=" singapore ", Currency=" usd ", Amount="$1,200.50", Department="it"), "Singapore", "2026-01"
    )

    assert result["IsValid"] is True
    assert result["Office"] == "Singapore" and result["Original_Office"] == " singapore "
    assert result["Currency"] == "USD" and result["Original_Currency"] == " usd "
    assert result["Department"] == "IT" and result["Original_Department"] == "it"
    assert result["Amount"] == 1200.50 and result["Original_Amount"] == "$1,200.50"


@pytest.mark.parametrize(
    ("field", "bad_value", "error"),
    [
        ("Amount", "USD 120.50", "Invalid Amount"),
        ("Amount", "(45.00)", "Invalid Amount"),
        ("Date", "01/15/2026", "Invalid Date"),
        ("Office", "Paris", "Invalid Office"),
        ("Department", "Legal", "Invalid Department"),
        ("ExpenseType", "Parking", "Invalid ExpenseType"),
        ("Currency", "EUR", "Unsupported Currency"),
    ],
)
def test_invalid_values_keep_original_evidence(field: str, bad_value: str, error: str) -> None:
    result = validate_and_normalize_row(row(**{field: bad_value}), "Singapore", "2026-01")

    assert error in str(result["ErrorType"])
    assert result[f"Original_{field}"] == bad_value


def test_unparseable_amount_is_blank_when_normalized_but_original_survives_in_workbook(tmp_path: Path) -> None:
    write_csv(tmp_path / "in" / "expense_tokyo_2026-01.csv", [row(Office="Tokyo", Amount="USD 120.50")])
    out = tmp_path / "out"

    run_report(tmp_path / "in", out)

    wb = load_workbook(out / "Expense_Report_2026-01.xlsx")
    for sheet in ("Raw_Data", "Error_Report"):
        (record,) = sheet_dicts(wb, sheet)
        assert record["Amount"] is None
        assert record["Original_Amount"] == "USD 120.50"
        assert record["ErrorType"] == "Invalid Amount"


def test_original_columns_present_in_raw_data_and_error_report(tmp_path: Path) -> None:
    write_csv(tmp_path / "in" / "expense_tokyo_2026-01.csv", [row(Office="Tokyo"), row(Office="Paris")])
    out = tmp_path / "out"

    run_report(tmp_path / "in", out)

    wb = load_workbook(out / "Expense_Report_2026-01.xlsx")
    for sheet in ("Raw_Data", "Error_Report"):
        headers = [c.value for c in wb[sheet][1]]
        assert all(column in headers for column in ORIGINAL_COLUMNS)


# --- file-level behaviour ----------------------------------------------------

def test_missing_required_columns_skips_file_and_logs_it(tmp_path: Path) -> None:
    bad_header = ",".join(c for c in REQUIRED_COLUMNS if c not in ("Amount", "Currency"))
    path = tmp_path / "in" / "expense_london_2026-01.csv"
    path.parent.mkdir()
    path.write_text(bad_header + "\n2026-01-01,London,IT,A,Meal,Lunch\n", encoding="utf-8")
    write_csv(tmp_path / "in" / "expense_tokyo_2026-01.csv", [row(Office="Tokyo")])
    out = tmp_path / "out"

    assert run_report(tmp_path / "in", out) == 0

    assert "expense_london_2026-01.csv: missing required columns: Amount, Currency" in log_text(out)
    m = metrics(load_workbook(out / "Expense_Report_2026-01.xlsx"))
    assert (m["Total Files Processed"], m["Total Records"]) == (1, 1)


@pytest.mark.parametrize(
    "name",
    ["expense_paris_2026-01.csv", "expense_tokyo_2026-1.csv", "expenses_tokyo_2026-01.csv", "tokyo.csv"],
)
def test_invalid_filename_is_rejected(name: str) -> None:
    assert parse_expense_filename(Path(name)) is None


def test_valid_filename_is_parsed_case_insensitively() -> None:
    parsed = parse_expense_filename(Path("Expense_New_York_2026-03.CSV".replace("New_York", "new_york")))

    assert parsed is not None and parsed.office == "New York" and parsed.report_month == "2026-03"


def test_invalid_filename_is_skipped_and_logged(tmp_path: Path) -> None:
    write_csv(tmp_path / "in" / "expense_paris_2026-01.csv", [row()])
    write_csv(tmp_path / "in" / "expense_tokyo_2026-01.csv", [row(Office="Tokyo")])
    out = tmp_path / "out"

    assert run_report(tmp_path / "in", out) == 0

    assert "Skipped expense_paris_2026-01.csv: invalid file name." in log_text(out)
    assert metrics(load_workbook(out / "Expense_Report_2026-01.xlsx"))["Total Files Processed"] == 1


def test_only_invalid_filenames_creates_no_workbook(tmp_path: Path) -> None:
    write_csv(tmp_path / "in" / "bogus.csv", [row()])
    out = tmp_path / "out"

    assert run_report(tmp_path / "in", out) == 0

    assert not list(out.glob("*.xlsx"))
    assert "Output File: N/A" in log_text(out)
    assert "Skipped bogus.csv: invalid file name." in log_text(out)


def test_multiple_report_months_abort_without_workbook(tmp_path: Path) -> None:
    write_csv(tmp_path / "in" / "expense_tokyo_2026-01.csv", [row(Office="Tokyo")])
    write_csv(tmp_path / "in" / "expense_london_2026-02.csv", [row(Office="London", Date="2026-02-02")])
    out = tmp_path / "out"

    assert run_report(tmp_path / "in", out) == 1

    assert not list(out.glob("*.xlsx"))
    assert "Aborted: multiple report months found (2026-01, 2026-02)." in log_text(out)


def test_no_csv_files_exits_cleanly_without_workbook(tmp_path: Path) -> None:
    (tmp_path / "in").mkdir()
    out = tmp_path / "out"

    assert run_report(tmp_path / "in", out) == 0

    assert not list(out.glob("*.xlsx"))
    assert "No CSV files found. No Excel report created." in log_text(out)


def test_missing_input_dir_is_treated_as_no_csv_files(tmp_path: Path) -> None:
    out = tmp_path / "out"

    assert run_report(tmp_path / "does_not_exist", out) == 0

    assert "No CSV files found" in log_text(out)


# --- workbook output ---------------------------------------------------------

def build_mixed_workbook(tmp_path: Path):
    write_csv(
        tmp_path / "in" / "expense_singapore_2026-01.csv",
        [
            row(Amount="100.10", Department="IT", ExpenseType="Meal"),
            row(Amount="$200.20", Department="Admin", ExpenseType="Travel"),
            row(Amount="bad"),
            row(Amount="10.00", Date="2026-02-01"),
            row(Amount="10.00", Office="Tokyo"),
        ],
    )
    write_csv(
        tmp_path / "in" / "expense_tokyo_2026-01.csv",
        [row(Office="tokyo", Amount="0.30", Department="IT", ExpenseType="Meal"), row(Office="Tokyo", Currency="EUR")],
    )
    out = tmp_path / "out"
    assert run_report(tmp_path / "in", out) == 0
    return load_workbook(out / "Expense_Report_2026-01.xlsx")


def test_workbook_is_created_with_expected_sheet_names(tmp_path: Path) -> None:
    wb = build_mixed_workbook(tmp_path)

    assert wb.sheetnames == EXPECTED_SHEETS
    assert (tmp_path / "out" / "process_log.txt").exists()


def test_totals_reconcile_after_filtering_invalid_rows(tmp_path: Path) -> None:
    wb = build_mixed_workbook(tmp_path)
    m = metrics(wb)

    assert m["Total Records"] == 7
    assert m["Valid Records"] + m["Error Records"] == m["Total Records"]
    assert (m["Valid Records"], m["Error Records"]) == (3, 4)
    assert m["Total Amount"] == pytest.approx(300.60)

    raw = sheet_dicts(wb, "Raw_Data")
    valid_sum = round(sum(r["Amount"] for r in raw if r["IsValid"]), 2)
    assert valid_sum == m["Total Amount"]
    assert len(sheet_dicts(wb, "Error_Report")) == m["Error Records"]
    for sheet in ("By_Office", "By_Department", "By_ExpenseType"):
        groups = sheet_dicts(wb, sheet)
        assert round(sum(g["TotalAmount"] for g in groups), 2) == m["Total Amount"]
        assert sum(g["RecordCount"] for g in groups) == m["Valid Records"]


def test_amount_columns_use_two_decimal_format_and_other_columns_do_not(tmp_path: Path) -> None:
    wb = build_mixed_workbook(tmp_path)

    raw = wb["Raw_Data"]
    headers = {c.value: c.column_letter for c in raw[1]}
    assert raw[f"{headers['Amount']}2"].number_format == "#,##0.00"
    assert raw[f"{headers['Original_Amount']}2"].number_format == "General"
    assert raw[f"{headers['RowNumber']}2"].number_format == "General"
    by_office = wb["By_Office"]
    assert by_office["B2"].number_format == "#,##0.00"
    assert by_office["C2"].number_format == "General"
    assert raw.freeze_panes == "A2" and raw["A1"].font.bold


# --- CSV text must never become a live formula in the workbook -------------------------

# Harmless probes: leading formula characters, whitespace that normalize_text() strips
# before the value reaches the normalized columns, and full-width look-alikes.
FORMULA_PAYLOADS = [
    "=1+1",
    "+1+1",
    "-1+1",
    "@SUM(1,1)",
    " =1+1",
    "\t=1+1",
    "\n=1+1",
    "\uff1d1+1",
    "\uff0b1+1",
    "\uff0d1+1",
    "\uff20SUM(1)",
]


def formula_cells(path: Path) -> list[tuple[str, str]]:
    workbook = load_workbook(path)
    return [
        (sheet.title, cell.coordinate)
        for sheet in workbook.worksheets
        for sheet_row in sheet.iter_rows()
        for cell in sheet_row
        if cell.data_type == "f"
    ]


def formula_elements_in_xml(path: Path) -> int:
    """Count <f> elements in the saved file itself, independent of openpyxl's reader."""
    with zipfile.ZipFile(path) as archive:
        return sum(
            len(re.findall(rb"<f[ >/]", archive.read(name)))
            for name in archive.namelist()
            if name.startswith("xl/worksheets/") and name.endswith(".xml")
        )


@pytest.mark.parametrize("payload", FORMULA_PAYLOADS, ids=lambda p: ascii(p))
def test_csv_text_never_becomes_a_formula_in_the_workbook(tmp_path: Path, payload: str) -> None:
    valid = row(Employee=payload, Description=payload)
    invalid = row(
        Date=payload,
        Office=payload,
        Department=payload,
        Employee=payload,
        ExpenseType=payload,
        Amount=payload,
        Currency=payload,
        Description=payload,
    )
    write_csv(tmp_path / "in" / "expense_singapore_2026-01.csv", [valid, invalid])
    out = tmp_path / "out"
    assert run_report(tmp_path / "in", out) == 0
    xlsx = out / "Expense_Report_2026-01.xlsx"

    assert formula_cells(xlsx) == []
    assert formula_elements_in_xml(xlsx) == 0

    wb = load_workbook(xlsx)
    assert wb.sheetnames == EXPECTED_SHEETS
    raw = sheet_dicts(wb, "Raw_Data")
    assert [r["IsValid"] for r in raw] == [True, False]
    # Original_* keeps the CSV text exactly, including leading whitespace.
    for record in (raw[1], sheet_dicts(wb, "Error_Report")[0]):
        for column in REQUIRED_COLUMNS:
            assert record[f"Original_{column}"] == payload
    assert raw[0]["Original_Employee"] == payload and raw[0]["Original_Description"] == payload
    # Normalized text is still the trimmed value, now stored as text.
    assert raw[0]["Employee"] == payload.strip()
    assert raw[0]["Description"] == payload.strip()
    # Numbers stay numbers and the valid-row total is unchanged.
    amount_col = [c.value for c in wb["Raw_Data"][1]].index("Amount") + 1
    assert wb["Raw_Data"].cell(row=2, column=amount_col).data_type == "n"
    assert metrics(wb)["Total Amount"] == pytest.approx(100.00)


def test_run_log_paths_starting_with_equals_are_not_formulas(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    write_csv(Path("=in") / "expense_singapore_2026-01.csv", [row()])
    assert run_report(Path("=in"), Path("+out")) == 0
    xlsx = Path("+out") / "Expense_Report_2026-01.xlsx"

    assert formula_cells(xlsx) == []
    assert formula_elements_in_xml(xlsx) == 0
    log = {r["Item"]: r["Value"] for r in sheet_dicts(load_workbook(xlsx), "Run_Log")}
    assert log["Input Folder"] == "=in"


def test_bundled_sample_data_still_reconciles_and_has_no_formulas(tmp_path: Path) -> None:
    assert run_report(PROJECT_ROOT / "input", tmp_path) == 0
    xlsx = tmp_path / "Expense_Report_2026-01.xlsx"
    wb = load_workbook(xlsx)

    assert wb.sheetnames == EXPECTED_SHEETS
    m = metrics(wb)
    assert (m["Total Files Processed"], m["Total Records"], m["Valid Records"], m["Error Records"]) == (3, 166, 150, 16)
    assert m["Total Amount"] == pytest.approx(75592.10)
    assert formula_cells(xlsx) == []
    assert formula_elements_in_xml(xlsx) == 0
